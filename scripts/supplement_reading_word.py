"""Add owner-supplied Word papers and workbook keys without removing old PDFs."""
import base64
import hashlib
import io
import json
import re
import secrets
import sys
import tomllib
import zipfile
from pathlib import Path

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(Path(__file__).parent))
from core import MAGIC,read_json,write_json


def cell_answer(value):
    if isinstance(value,bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value,(int,float)):
        return str(int(value)) if int(value)==value else str(value)
    return str(value).strip() if value is not None else ''


def workbook_keys(data,issues=None):
    issues=issues if issues is not None else []
    result={}
    workbook=load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    try:
        for sheet in workbook:
            match=re.fullmatch(r'\s*Test\s*(\d+)\s*',sheet.title,re.I)
            if not match:
                continue
            answers={}
            conflicts=set()
            rows=list(sheet.iter_rows(max_row=200,max_col=26,values_only=True))
            for col in range(0,26,2):
                records=[]
                for row_index,row in enumerate(rows,1):
                    number,value=row[col:col+2]
                    if isinstance(number,bool) or not isinstance(number,(int,float)) or int(number)!=number or not 1<=number<=40:
                        continue
                    value=cell_answer(value)
                    if value:
                        records.append((row_index,int(number),value))
                previous=0
                for i,(row_index,number,value) in enumerate(records):
                    following=records[i+1][1] if i+1<len(records) else None
                    if number<=previous or (previous>0 and following is not None and previous<following<number):
                        issues.append({'sheet':sheet.title,'row':row_index,'column':col+1,'number':number,'reason':'Question number out of sequence; not imported'})
                        continue
                    previous=number
                    number=str(number)
                    parsed=[part.strip() for part in re.split(r'\s*(?:/|;|\bor\b)\s*',value,flags=re.I) if part.strip()]
                    if number in answers and answers[number]!=parsed:
                        issues.append({'sheet':sheet.title,'row':row_index,'column':col+1,'number':int(number),'reason':'Conflicting duplicate answer; not imported'})
                        conflicts.add(number)
                        answers.pop(number,None)
                    elif number not in conflicts:
                        answers[number]=parsed
            if answers:
                result[int(match[1])]=answers
    finally:
        workbook.close()
    return result


def merge(archives):
    catalog=read_json(ROOT/'data/catalog.json',{})
    sources=read_json(ROOT/'data/local_sources.json',{})
    resources=catalog['resources']
    tests={t['id']:t for t in catalog['tests']}
    changed=[]
    report=[]
    for volume,path in archives:
        source=f'reading-word-v{volume}'
        sources[source]=str(path.resolve())
        with zipfile.ZipFile(path) as archive:
            key_name=next(n for n in archive.namelist() if n.lower().endswith('.xlsx'))
            key_id=hashlib.sha256((source+':'+key_name).encode()).hexdigest()[:16]
            resources[key_id]={'id':key_id,'source':source,'entry':key_name,'name':Path(key_name).name,'ext':'.xlsx','size':archive.getinfo(key_name).file_size,'role':'key','volume':volume}
            issues=[]
            keys=workbook_keys(archive.read(key_name),issues)
            found=set()
            for item in archive.infolist():
                if not item.filename.lower().endswith('.docx'):
                    continue
                match=re.search(r'\btest\s*(\d+)',Path(item.filename).name,re.I)
                if not match:
                    raise ValueError('Cannot identify test: '+item.filename)
                number=int(match[1])
                if number in found or not 1<=number<=10:
                    raise ValueError('Duplicate or unexpected test number')
                found.add(number)
                tid=f'reading-v{volume}-t{number}'
                test=tests[tid]
                rid=hashlib.sha256((source+':'+item.filename).encode()).hexdigest()[:16]
                resources[rid]={'id':rid,'source':source,'entry':item.filename,'name':Path(item.filename).name,'ext':'.docx','size':item.file_size,'role':'paper','volume':volume}
                old_ids=list(test['resources'])
                test['resources']=[rid]+[r for r in old_ids if r!=rid]
                test['preferred_paper']=rid
                if number in keys:
                    test.update(answers=keys[number],groups=[],key_source=key_id,key_count=len(keys[number]))
                    test['key_import_notes']=[note for note in issues if re.fullmatch(r'\s*Test\s*'+str(number)+r'\s*',note['sheet'],re.I)]
                    if key_id not in test['resources']:
                        test['resources'].append(key_id)
                changed.append(tid)
                report.append({'test':tid,'word':item.filename,'answers':test['key_count'],'pdf_preserved':sum(resources[r]['ext']=='.pdf' and resources[r]['role']=='paper' for r in old_ids),'key_notes':test.get('key_import_notes',[])})
            if found!=set(range(1,11)):
                raise ValueError(f'VOL {volume}: expected tests 1–10, found {sorted(found)}')
    key=base64.urlsafe_b64decode(tomllib.loads((ROOT/'.streamlit/secrets.toml').read_text(encoding='utf-8'))['DATA_KEY'])
    handles={s:zipfile.ZipFile(p) for s,p in sources.items()}
    try:
        for tid in changed:
            test=tests[tid]
            buffer=io.BytesIO()
            with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED,compresslevel=2) as pack:
                for rid in test['resources']:
                    resource=resources[rid]
                    pack.writestr(rid+resource['ext'],handles[resource['source']].read(resource['entry']))
            nonce=secrets.token_bytes(12)
            payload=MAGIC+nonce+AESGCM(key).encrypt(nonce,buffer.getvalue(),MAGIC)
            test['pack']=tid+'-'+hashlib.sha256(payload).hexdigest()[:12]+'.bingo'
            (ROOT/'data/packs'/test['pack']).write_bytes(payload)
    finally:
        for archive in handles.values():
            archive.close()
    write_json(ROOT/'data/catalog.json',catalog)
    # Keep the tracked flat deployment mirror in sync, without touching copies.
    if (ROOT/'catalog.json').exists():
        write_json(ROOT/'catalog.json',catalog)
    write_json(ROOT/'data/local_sources.json',sources)
    write_json(ROOT/'.bingo/reading-word-report.json',report)
    print(json.dumps({'word_papers':len(changed),'total_tests':len(tests),'full_keys':sum(t['key_count']==40 for t in tests.values()),'details':report},ensure_ascii=False),flush=True)


if __name__=='__main__':
    merge([(3,ROOT.parent/'VOL3-READING-WORD.zip'),(4,ROOT.parent/'VOL4-READING-WORD.zip'),(5,ROOT.parent/'vol5-reading-word.zip')])
