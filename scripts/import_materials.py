"""Index the owner's ZIPs without extracting multi-gigabyte media into Git."""
import argparse
import hashlib
import io
import json
import re
import sys
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]


def plain(value):
    value = unicodedata.normalize('NFKD', str(value)).replace('đ', 'd').replace('Đ', 'D')
    return ''.join(c for c in value if not unicodedata.combining(c)).lower()


def variants(value):
    value = re.sub(r'\s+', ' ', str(value)).strip(' .;:')
    if not value or len(value) > 100:
        return []
    # Letter keys often carry a description that isn't part of the answer.
    if re.match(r'^[A-H]\s+[^/,]', value) and not re.match(r'^[A-H]\s*(?:and|or)\b', value):
        value = value[0]
    value = re.sub(r'\s*\([^)]*\)\s*$', '', value).strip()
    out = re.split(r'\s*(?:/|\bor\b|;|\|)\s*', value, flags=re.I)
    if len(out) == 1 and value.count(',') >= 2 and not re.search(r'\d,\d{3}', value):
        out = value.split(',')
    return [v.strip(' .') for v in out if v.strip(' .')]


def parse_key(text):
    answers = {}
    groups = []
    # Explicit Vietnamese question/answer blocks take priority over incidental numbers.
    blocks = list(re.finditer(r'(?:Câu\s*(?:hỏi|số)?\s*)(\d{1,2})(?:\s*[-–]\s*(\d{1,2}))?\s*[:.)]?', text, re.I))
    for i, m in enumerate(blocks):
        n, end = int(m[1]), int(m[2] or m[1])
        if not 1 <= n <= end <= 40:
            continue
        body = text[m.end():blocks[i+1].start() if i+1 < len(blocks) else len(text)]
        explicit = re.search(r'Đáp án\s*:\s*([^\n]+)', body, re.I)
        if explicit:
            value = explicit[1]
        elif re.match(r'\s*:\s*\n', text[m.start():]):
            value = body.strip().split('\n')[0]
        elif 'cau so' in plain(m[0]):
            value = body.strip().split('\n')[0] if body.strip() else ''
        else:
            continue
        if end > n:
            parts = re.findall(r'\b[A-H]\b', value)
            if len(parts) == end-n+1:
                groups.append({'numbers': list(range(n, end+1)), 'answers': parts})
                for k, v in zip(range(n, end+1), parts):
                    answers[str(k)] = [v]
        elif variants(value):
            answers[str(n)] = variants(value)
    if len(answers) >= 30:
        return answers, groups
    # Numbered short keys, e.g. "1. city" or "1 C: quoted evidence".
    for line in text.splitlines():
        line = line.strip()
        m = re.match(r'^(\d{1,2})\s*(?:[.)]\s*|\s+)([^\n]+)$', line)
        if not m or not 1 <= int(m[1]) <= 40:
            continue
        value = m[2].strip()
        if ':' in value:
            value = value.split(':', 1)[0]
        if len(value.split()) <= 5 and len(value) <= 60 and not re.search(r'_{2,}|\.{3,}|\b(?:questions?|paragraph|passage|section)\b', value, re.I):
            vals = variants(value)
            if vals:
                answers.setdefault(str(int(m[1])), vals)
    return answers, groups


def text_of(z, name):
    data = io.BytesIO(z.read(name))
    if name.lower().endswith('.pdf'):
        return '\n'.join(p.extract_text() or '' for p in PdfReader(data).pages)
    doc = Document(data)
    return '\n'.join([p.text for p in doc.paragraphs] + ['\n'.join('  '.join(c.text for c in r.cells) for r in t.rows) for t in doc.tables])


def build(paths):
    resources = {}
    tests = {}
    sources = {}
    issues = []
    key_candidates = defaultdict(list)
    for path in paths:
        kind = 'listening' if 'listening' in path.name.lower() else 'reading'
        sources[kind] = str(path.resolve())
        with zipfile.ZipFile(path) as z:
            for entry in z.infolist():
                name = entry.filename
                ext = Path(name).suffix.lower()
                if ext not in {'.pdf', '.docx', '.mp3', '.mp4', '.m4a', '.wma', '.jpg', '.png', '.xlsx'}:
                    continue
                vm = re.search(r'vol\s*(\d+)', name, re.I)
                if not vm:
                    continue
                vol = int(vm[1])
                p = plain(name)
                nums = re.findall(r'(?:test|tets)\s*(\d+)', p)
                if not nums and '/explain reading/' in p:
                    nums = [str(int(Path(name).stem))] if Path(name).stem.isdigit() else []
                if not nums:
                    # Keep unmatched reference files accessible in a separate library.
                    num = None
                else:
                    num = int(nums[-1])
                rid = hashlib.sha256((kind + ':' + name).encode()).hexdigest()[:16]
                role = 'paper'
                if ext in {'.mp3', '.mp4', '.m4a', '.wma'}:
                    role = 'audio'
                elif ext == '.xlsx' or any(w in p for w in ['dap an', '/key/', 'key ', 'key-', 'explain', 'explanation', 'explantion', 'giai thich']):
                    role = 'key'
                elif 'transcript' in p or 'transcrip' in p:
                    role = 'transcript'
                if '/bao song ngu/' in p:
                    role = 'reference'
                    num = None
                resources[rid] = {'id': rid, 'source': kind, 'entry': name, 'name': Path(name).name, 'ext': ext, 'size': entry.file_size, 'role': role, 'volume': vol}
                if num:
                    tid = f'{kind}-v{vol}-t{num}'
                    test = tests.setdefault(tid, {'id': tid, 'skill': kind, 'volume': vol, 'number': num, 'title': f'VOL {vol} · Test {num}', 'resources': [], 'answers': {}, 'groups': [], 'key_source': None})
                    test['resources'].append(rid)
                    if role == 'key' and ext in {'.docx', '.pdf'}:
                        try:
                            a, g = parse_key(text_of(z, name))
                            if a:
                                key_candidates[tid].append((a, g, rid))
                        except Exception as e:
                            issues.append({'resource': rid, 'issue': type(e).__name__})
                if ext == '.xlsx':
                    workbook = load_workbook(io.BytesIO(z.read(name)), data_only=True)
                    for sheet in workbook:
                        match = re.search(r'test\s*(\d+)', sheet.title, re.I)
                        if not match:
                            continue
                        tid = f'{kind}-v{vol}-t{int(match[1])}'
                        a = {}
                        for row in sheet.values:
                            for col in range(len(row)-1):
                                n, value = row[col:col+2]
                                if isinstance(n, (int, float)) and int(n) == n and 1 <= n <= 40 and isinstance(value, str):
                                    vals = variants(value)
                                    if vals:
                                        a[str(int(n))] = vals
                        if a:
                            key_candidates[tid].append((a, [], rid))
    # Only retain tests with an actual question paper, not orphan keys/audio.
    tests = {k:v for k,v in tests.items() if any(resources[r]['role'] == 'paper' for r in v['resources'])}
    for tid, test in tests.items():
        choices = key_candidates.get(tid, [])
        if choices:
            a, g, rid = max(choices, key=lambda item: (len(item[0]), resources[item[2]]['ext'] == '.xlsx'))
            test.update(answers=a, groups=g, key_source=rid)
            if rid not in test['resources']:
                test['resources'].append(rid)
        test['resources'].sort(key=lambda r: (resources[r]['role'], plain(resources[r]['name'])))
        test['key_count'] = len(test['answers'])
        test['has_audio'] = any(resources[r]['role'] == 'audio' for r in test['resources'])
    result = {'version': 1, 'tests': sorted(tests.values(), key=lambda t:(t['skill'],t['volume'],t['number'])), 'resources': resources, 'issues': issues}
    (ROOT/'data/catalog.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    (ROOT/'data/local_sources.json').write_text(json.dumps(sources, indent=2), encoding='utf-8')
    print(json.dumps({'tests':len(tests), 'reading':sum(t['skill']=='reading' for t in tests.values()), 'listening':sum(t['skill']=='listening' for t in tests.values()), 'full_keys':sum(t['key_count']==40 for t in tests.values()), 'audio':sum(t['has_audio'] for t in tests.values()), 'issues':len(issues)}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('archives', nargs='+', type=Path)
    build(parser.parse_args().archives)
