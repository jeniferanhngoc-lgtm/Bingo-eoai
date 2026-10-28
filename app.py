"""Bingo: personal IELTS Listening and Reading, deployable on Streamlit Cloud."""
import hashlib
import base64
import hmac
import io
import json
import os
import time
from datetime import datetime, timezone
from pathlib import Path

import streamlit as st

from core import ROOT, audio_files, audio_part, clean_answers, missing_audio_parts, band, docx_html, grade, read_json, resource_bytes, write_json

st.set_page_config(page_title='Bingo · IELTS Practice', page_icon='🎯', layout='wide')
css_path = ROOT/'assets/bingo.css' if (ROOT/'assets/bingo.css').exists() else ROOT/'bingo.css'
css = css_path.read_text(encoding='utf-8')
for asset,mime in [('bg-header.jpg','image/jpeg'),('bg-body.png','image/png'),('bg-footer.jpg','image/jpeg'),('bg-menu.png','image/png')]:
    path=ROOT/'assets/dreamy'/asset
    css=css.replace('__'+asset+'__','data:'+mime+';base64,'+base64.b64encode(path.read_bytes()).decode())
st.html('<style>'+'\n'.join(line for line in css.splitlines() if not line.startswith('@import'))+'</style>')
st.html('<header class="dreamy-header"><div class="dreamy-brand">Bingo</div><div class="dreamy-tagline">A little practice. A brighter day.</div></header>')


def setting(name, default=''):
    try:
        return st.secrets.get(name, os.getenv(name, default))
    except FileNotFoundError:
        return os.getenv(name, default)


def settings(tid=None):
    config={key:setting(key) for key in ['DATA_REPOSITORY','DATA_RELEASE','GITHUB_TOKEN','DATA_KEY']}
    if tid:
        config['PACK_NAME']=by_id[tid].get('pack','')
    return config


def gate():
    password = setting('BINGO_PASSWORD')
    if not password:
        st.title('Bingo · Thiết lập mật khẩu')
        st.info('Thêm BINGO_PASSWORD vào Streamlit Secrets để mở ứng dụng. Bingo luôn yêu cầu mật khẩu, kể cả khi mã nguồn công khai.')
        st.code('BINGO_PASSWORD = "mat-khau-cua-ban"', language='toml')
        st.stop()
    fingerprint = hashlib.sha256(password.encode()).hexdigest()
    if st.session_state.get('authenticated') == fingerprint:
        return
    _, center, _ = st.columns([1,1.2,1])
    with center:
        st.html('<div class="login-intro"><h2>Chào bạn,</h2><p>Một chút luyện tập cho ngày mới. Nhập mật khẩu để vào không gian IELTS của riêng bạn.</p></div>')
        with st.form('login'):
            value = st.text_input('Mật khẩu của bạn', type='password')
            submit = st.form_submit_button('Vào Bingo →', type='primary', width='stretch')
        if submit:
            if hmac.compare_digest(value.encode(), password.encode()):
                st.session_state.authenticated = fingerprint
                st.rerun()
            else:
                st.error('Mật khẩu chưa đúng.')
    st.stop()


gate()
catalog_path = ROOT/'data/catalog.json' if (ROOT/'data/catalog.json').exists() else ROOT/'catalog.json'
catalog = read_json(catalog_path, {'tests':[], 'resources':{}})
tests = catalog['tests']
by_id = {t['id']:t for t in tests}
resources = catalog['resources']
STATE = ROOT/'.bingo/progress.json'
if 'progress' not in st.session_state:
    st.session_state.progress = read_json(STATE, {'drafts':{},'history':[], 'favorites':[]})
progress = st.session_state.progress


def save():
    write_json(STATE, progress)


def draft(test):
    d = progress['drafts'].setdefault(test['id'], {'answers':{},'remaining':3600 if test['skill']=='reading' else 2400,'deadline':None})
    d.setdefault('created',time.time_ns())
    return d


def record(test):
    d = draft(test)
    if d.get('deadline'):
        d['remaining'] = max(0,int(d['deadline']-time.time()))
    d['deadline'] = None
    marked = grade(test,d['answers'])
    item = {'id':f"{test['id']}-{time.time_ns()}", 'test_id':test['id'], 'title':test['title'], 'skill':test['skill'], 'date':datetime.now(timezone.utc).isoformat(), 'responses':dict(d['answers']), 'results':marked, 'score':sum(r['correct'] for r in marked.values()), 'total':len(marked)}
    progress['history'].insert(0,item)
    save()
    st.session_state.result_id = item['id']


@st.cache_data(max_entries=4, show_spinner=False)
def load_resource(rid, tid, config):
    return resource_bytes(resources[rid], tid, config)


@st.cache_data(max_entries=4, show_spinner=False)
def load_word_html(rid, tid, config):
    return docx_html(load_resource(rid, tid, config))


@st.fragment
def viewer(rid, tid):
    r = resources[rid]
    try:
        with st.spinner('Đang mở tài liệu…'):
            data = load_resource(rid,tid,settings(tid))
        if r['ext']=='.pdf':
            from reader import highlight_reader
            highlight_reader(data,'.pdf',rid,progress,save)
        elif r['ext']=='.docx':
            from reader import highlight_reader
            highlight_reader(data,'.docx',rid,progress,save,word_html=docx_html(data))
        elif r['ext'] in {'.jpg','.png'}:
            st.image(data)
        elif r['ext']=='.xlsx':
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(data),data_only=True)
            sheet_name = st.selectbox('Bảng đáp án',book.sheetnames,index=next((i for i,n in enumerate(book.sheetnames) if n.lower()==f"test {by_id[tid]['number']}"),0),key='sheet-'+rid)
            rows = list(book[sheet_name].values)
            st.dataframe([[str(v or '') for v in row] for row in rows],hide_index=True,width='stretch')
        st.download_button('Tải tài liệu gốc ↗', data,file_name=r['name'],key='download-'+rid)
    except Exception as e:
        # Don't leak private URLs, tokens or exception internals to the page.
        st.error('Chưa mở được tài liệu. Kiểm tra cấu hình dữ liệu hoặc kết nối mạng trong mục Hướng dẫn.')
        if isinstance(e,FileNotFoundError):
            st.caption(str(e))


def open_test(tid):
    st.session_state.active = tid
    st.session_state.pop('result_id',None)


with st.container(key='dreamy-nav'):
    nav,logout=st.columns([6,1])
    page = nav.radio('Điều hướng',['Kho đề','Lịch sử','Sao lưu & hướng dẫn'],label_visibility='collapsed',horizontal=True)
    if logout.button('Đăng xuất',width='stretch'):
        st.session_state.clear()
        st.rerun()


def library():
    st.html('<section class="hero"><div class="eyebrow">LISTENING & READING / YOUR STUDY JOURNAL</div><h1>Mỗi ngày, một bước tiến.</h1><p>Chọn một bộ đề, dành thời gian cho mình và bắt đầu. Tất cả VOL 1–8 đều ở đây.</p></section>')
    a,b,c = st.columns(3)
    a.metric('Bài Reading',sum(t['skill']=='reading' for t in tests))
    b.metric('Bài Listening',sum(t['skill']=='listening' for t in tests))
    c.metric('Lượt đã luyện',len(progress['history']))
    st.write('')
    a,b,c,d = st.columns([1.2,1,2,1])
    skill = a.selectbox('Kỹ năng',['Tất cả','Reading','Listening'])
    volumes = sorted({t['volume'] for t in tests})
    vol = b.selectbox('Bộ đề · VOL 1–8',['Tất cả']+[f'VOL {n}' for n in volumes])
    search = c.text_input('Tìm bài',placeholder='Ví dụ: Test 3')
    favorite_only = d.checkbox('Đã lưu',value=False)
    shown = [t for t in tests if (skill=='Tất cả' or t['skill']==skill.lower()) and (vol=='Tất cả' or t['volume']==int(vol.split()[1])) and (not search or search.lower() in t['title'].lower()) and (not favorite_only or t['id'] in progress['favorites'])]
    st.caption(f'{len(shown)} bài phù hợp · Đề gốc và đáp án được giữ nguyên để đối chiếu.')
    page_count = max(1,(len(shown)+11)//12)
    st.caption(f'Trang 1–{page_count} · Mỗi trang có tối đa 12 bài. Chọn VOL ở trên để đến thẳng bộ đề bạn cần.')
    number = st.number_input(f'Chọn trang (1–{page_count})',min_value=1,max_value=page_count,value=1,step=1)
    st.caption(f'Đang xem trang {number}/{page_count} · Bài {(number-1)*12+1 if shown else 0}–{min(number*12,len(shown))} trên {len(shown)} bài')
    for offset in range(0,len(shown[(number-1)*12:number*12]),3):
        for column,t in zip(st.columns(3),shown[(number-1)*12+offset:(number-1)*12+offset+3]):
            with column,st.container(border=True):
                st.html(f'<span class="tag {t["skill"]}">{t["skill"].upper()}</span><div class="bingo-card-title">{t["title"]}</div>')
                st.caption(('60 phút' if t['skill']=='reading' else '40 phút')+' · Phiếu 40 câu')
                st.caption(f"Chấm tự động {t['key_count']}/40 câu" if t['key_count'] else 'Đối chiếu đáp án gốc')
                if t['skill']=='listening' and not t['has_audio']:
                    st.warning('ZIP chưa có audio cho bài này.',icon='🎧')
                elif t['skill']=='listening' and missing_audio_parts(t,resources):
                    st.caption('⚠ Thiếu audio phần '+', '.join(map(str,missing_audio_parts(t,resources)))+' trong ZIP')
                done = sum(h['test_id']==t['id'] for h in progress['history'])
                if done:
                    st.caption(f'✓ Đã luyện {done} lần')
                x,y = st.columns([3,1])
                x.button('Tiếp tục →' if t['id'] in progress['drafts'] else 'Luyện ngay →',key='open-'+t['id'],type='primary',width='stretch',on_click=open_test,args=(t['id'],))
                star = t['id'] in progress['favorites']
                if y.button('★' if star else '☆',key='star-'+t['id'],help='Lưu bài'):
                    if star:
                        progress['favorites'].remove(t['id'])
                    else:
                        progress['favorites'].append(t['id'])
                    save()
                    st.rerun()


@st.fragment
def audio_panel(test):
    audio = audio_files(test,resources)
    if not audio:
        st.warning('Chưa có audio trong ZIP cho bài này. Bạn có thể luyện bài khác hoặc bổ sung audio sau.')
        return
    missing = missing_audio_parts(test,resources)
    with st.container(key='exam-audio'):
        pick,player = st.columns([1,2],vertical_alignment='center')
        rid = pick.selectbox('Bản ghi',audio,format_func=lambda r:(f'Phần {audio_part(resources[r])} · ' if audio_part(resources[r]) else '')+resources[r]['name'],key='audio-'+test['id'],label_visibility='collapsed')
        try:
            data = load_resource(rid,test['id'],settings(test['id']))
            ext = resources[rid]['ext']
            with player:
                if ext == '.wma':
                    st.download_button('Tải audio WMA (trình duyệt có thể không phát được)',data,file_name=resources[rid]['name'])
                else:
                    st.audio(data,format={'.mp3':'audio/mpeg','.m4a':'audio/mp4','.mp4':'audio/mp4'}.get(ext,'audio/mpeg'))
        except Exception:
            player.error('Chưa tải được audio. Kiểm tra kho dữ liệu và kết nối mạng.')
        if missing:
            st.caption('ZIP chưa có audio phần '+', '.join(map(str,missing))+'. Bạn có thể luyện các phần đang có.')


def reset_exam_view(tid):
    for k in [k for k in st.session_state if k==f'exam-frozen-{tid}' or k==f'exam-event-{tid}']:
        del st.session_state[k]


def exam_state(test,rid,page,view,session):
    """Arguments for the exam component. They stay identical between reruns, because Streamlit
    derives the component's identity from them and would reload the page whenever they change."""
    tid=test['id']
    request=(rid,page,view)
    frozen=st.session_state.get('exam-frozen-'+tid)
    if frozen and frozen['session']==session and frozen['request']==request:
        return frozen
    d=draft(test)
    now=time.time()
    if d.get('deadline'):
        remaining,running=max(0,int(d['deadline']-now)),True
    else:
        remaining,running=int(d['remaining']),not d.get('paused')
        if running and remaining>0:
            d['deadline']=now+remaining
            save()
    try:
        ext=resources[rid]['ext']
        data=load_resource(rid,tid,settings(tid))
        from reader import prepare_document
        doc,stored,section=prepare_document(data,ext,rid,progress,page,view,load_word_html(rid,tid,settings(tid)) if ext=='.docx' else None)
    except Exception as e:
        message=str(e) if isinstance(e,FileNotFoundError) else 'Chưa mở được tài liệu. Kiểm tra cấu hình dữ liệu hoặc kết nối mạng trong mục Hướng dẫn.'
        doc={'kind':'empty','message':message,'page':1,'pages':0,'can_text':False,'view':'image','document_key':rid+':error'}
        stored={'marks':[],'updated':0}
    frozen={'session':session,'request':request,'doc':doc,'annotations':stored,'answers':dict(d['answers']),'answers_updated':d.get('updated',0),'remaining':remaining,'running':running,'stamp':int(now*1000)}
    st.session_state['exam-frozen-'+tid]=frozen
    return frozen


def apply_exam_event(test,event,rid,doc):
    tid=test['id']
    d=draft(test)
    try:
        d['answers']=clean_answers(event.get('answers',{}))
    except ValueError:
        st.warning('Phiếu trả lời chưa hợp lệ nên chưa được lưu.')
        return
    now=time.time()
    d['updated']=int(now*1000)
    remaining=event.get('remaining')
    if isinstance(remaining,(int,float)) and not isinstance(remaining,bool) and remaining==remaining:
        remaining=max(0,min(int(remaining),7200))
        d['remaining']=remaining
        if event.get('running') and event.get('type') not in {'back'}:
            d['deadline']=now+remaining
            d['paused']=False
        else:
            d['deadline']=None
            d['paused']=event.get('type')!='back'
    kind=event.get('type')
    if kind=='highlight':
        from reader import store_highlight
        if event.get('document_key')==doc['document_key'] and doc['kind'] in {'image','text'}:
            section=doc['document_key'].split(':',1)[1]
            try:
                store_highlight(progress,rid,section,doc['kind'],(event.get('annotations') or {}))
            except ValueError:
                st.warning('Highlight chưa hợp lệ nên chưa được lưu.')
    save()
    if kind in {'submit','solution'}:
        record(test)
        reset_exam_view(tid)
        st.rerun()
    elif kind=='back':
        reset_exam_view(tid)
        st.session_state.pop('active',None)
        st.session_state.pop('result_id',None)
        st.session_state.pop('show_review',None)
        st.rerun()
    elif kind=='page' and isinstance(event.get('page'),int):
        st.session_state['exam-page-'+rid]=max(1,event['page'])
        st.rerun(scope='fragment')
    elif kind=='view' and event.get('view') in {'image','text'}:
        st.session_state['exam-view-'+rid]=event['view']
        st.rerun(scope='fragment')
    elif kind=='paper' and event.get('paper') in [r for r in test['resources'] if resources[r]['role']=='paper']:
        st.session_state['paper-'+tid]=event['paper']
        st.rerun(scope='fragment')


@st.fragment
def exam_screen(test):
    from reader import exam_component
    tid=test['id']
    d=draft(test)
    papers=[r for r in test['resources'] if resources[r]['role']=='paper']
    rid=st.session_state.get('paper-'+tid)
    if rid not in papers:
        rid=papers[0]
    session=f"{tid}:{d['created']}"
    frozen=exam_state(test,rid,int(st.session_state.get('exam-page-'+rid,1)),st.session_state.get('exam-view-'+rid,'image'),session)
    event=exam_component(
        session=session,
        created=str(d['created']),
        test={'id':tid,'title':test['title'],'skill':test['skill']},
        papers=[{'id':r,'name':resources[r]['name']} for r in papers],
        paper=rid,
        doc=frozen['doc'],
        annotations=frozen['annotations'],
        answers=frozen['answers'],
        answers_updated=frozen['answers_updated'],
        remaining=frozen['remaining'],
        running=frozen['running'],
        stamp=frozen['stamp'],
        key='exam-'+tid,
        default=None,
    )
    if isinstance(event,dict) and event.get('session')==session and event.get('nonce')!=st.session_state.get('exam-event-'+tid):
        st.session_state['exam-event-'+tid]=event['nonce']
        apply_exam_event(test,event,rid,frozen['doc'])


def result_view(item,test):
    st.success('Đã lưu lượt luyện của bạn.')
    a,b,c=st.columns(3)
    a.metric('Câu đúng',f"{item['score']}/{item['total']}" if item['total'] else 'Chưa có khóa chấm')
    b.metric('Đã trả lời',f"{sum(bool(v.strip()) for v in item['responses'].values())}/40")
    c.metric('Band tham khảo',f"{band(item['score'],test['skill']):.1f}" if item['total']==40 else '—')
    st.caption('Đáp án được nhập tự động từ tài liệu bạn cung cấp. Hãy đối chiếu bản gốc khi có khác biệt; band chỉ để tham khảo. Reading dùng bảng quy đổi Academic.')
    if item['total']<40:
        st.info(f"Bài này mới có khóa chấm cho {item['total']}/40 câu. Các câu còn lại cần đối chiếu tài liệu gốc; chưa tính band toàn bài.")
    rows=[]
    for n in range(1,41):
        r=item['results'].get(str(n))
        rows.append({'Câu':n,'Bạn trả lời':item['responses'].get(str(n),''),'Đáp án':r['expected'] if r else 'Đối chiếu tài liệu gốc','Kết quả':('✓ Đúng' if r['correct'] else '✗ Chưa đúng') if r else 'Chưa chấm'})
    st.dataframe(rows,hide_index=True,width='stretch',height=500)
    keys=[r for r in test['resources'] if resources[r]['role'] in {'key','transcript'}]
    if keys:
        rid=st.selectbox('Đáp án / lời giải / transcript gốc',keys,format_func=lambda r:resources[r]['name'])
        viewer(rid,test['id'])
    if st.button('Làm lại bài này',type='primary'):
        progress['drafts'].pop(test['id'],None)
        for k in list(st.session_state):
            if k.startswith(f'ans-{test["id"]}-'):
                del st.session_state[k]
        st.session_state.pop('result_id',None)
        save()
        st.rerun()


EXAM_CSS = """/* Exam mode: the window is the exam room. Only the two panes inside the exam component scroll. */
html,body,.stApp,[data-testid="stApp"],[data-testid="stAppViewContainer"],[data-testid="stMain"],section.main,.stMain{overflow:hidden!important;height:100vh!important;max-height:100vh!important;background:#fff!important}
[data-testid="stHeader"],[data-testid="stToolbar"],[data-testid="stDecoration"],[data-testid="stStatusWidget"],.dreamy-header,.st-key-dreamy-nav,.footer{display:none!important}
.block-container,[data-testid="stMainBlockContainer"]{max-width:100%!important;width:100%!important;padding:0!important;margin:0!important;min-height:0!important;box-shadow:none!important;background:#fff!important}
[data-testid="stMain"]{background:#fff!important}
[data-testid="stMainBlockContainer"] [data-testid="stVerticalBlock"],.block-container [data-testid="stVerticalBlock"]{gap:0!important}
/* never fade the exam while Streamlit is working */
[data-stale="true"],[data-testid="stElementContainer"][data-stale="true"]{opacity:1!important;transition:none!important}
[data-testid="stElementContainer"]:has(iframe){margin:0!important;line-height:0}
iframe{border:0!important;display:block}
/* Listening: compact audio strip above the exam */
.st-key-exam-audio{background:#f1f1f1;border-bottom:1px solid #ddd;padding:6px 14px 4px}
.st-key-exam-audio [data-testid="stHorizontalBlock"]{gap:12px!important;align-items:center}
.st-key-exam-audio [data-testid="stSelectbox"] label{display:none}
.st-key-exam-audio audio{width:100%;height:40px}
.st-key-exam-audio [data-testid="stAudio"],.st-key-exam-audio [data-testid="stAudio"] audio{margin:0}
.st-key-exam-audio [data-testid="stCaptionContainer"]{margin:0}
"""


def exam_css():
    return '<style>'+EXAM_CSS+'</style>'


def practice(test):
    item=next((h for h in progress['history'] if h['id']==st.session_state.get('result_id')),None)
    if item:
        if st.button('← Về kho đề'):
            st.session_state.pop('active',None)
            st.session_state.pop('result_id',None)
            st.session_state.pop('show_review',None)
            st.rerun()
        st.html(f'<div class="eyebrow">{test["skill"].upper()} / YOUR PRACTICE</div>')
        st.title(test['title'])
        result_view(item,test)
        return
    # Exam mode: the browser window is the exam room. Only the two panes scroll, never the page.
    st.html(exam_css())
    if test['skill']=='listening':
        audio_panel(test)
    exam_screen(test)


def history():
    st.title('Nhìn lại hành trình')
    st.caption('Mỗi lượt luyện là một bước tiến của bạn.')
    if not progress['history']:
        st.info('Chưa có lượt luyện. Chọn một bài trong Kho đề để bắt đầu.')
    for h in progress['history']:
        with st.container(border=True):
            a,b=st.columns([4,1])
            a.subheader(f"{h['skill'].title()} · {h['title']}")
            a.caption(f"{h['date'][:10]} · {h['score']}/{h['total']} câu đã chấm")
            if b.button('Xem lại →',key=h['id']):
                st.session_state.active=h['test_id']
                st.session_state.result_id=h['id']
                st.session_state.show_review=True
                st.rerun()


def backup():
    st.title('Sao lưu & hướng dẫn')
    st.write('Bingo lưu câu trả lời khi bạn rời ô nhập. Trên máy cá nhân, tiến độ nằm trong thư mục .bingo. Trên Streamlit Cloud, dữ liệu trên ổ đĩa có thể mất khi ứng dụng khởi động lại; hãy tải bản sao lưu sau mỗi buổi học.')
    st.download_button('↓ Tải bản sao lưu tiến độ',json.dumps(progress,ensure_ascii=False,indent=2).encode('utf-8'),file_name='bingo-progress.json',mime='application/json',type='primary')
    uploaded=st.file_uploader('Khôi phục tiến độ từ Bingo',type=['json'])
    if uploaded and st.button('Khôi phục bản sao lưu'):
        try:
            if uploaded.size>5_000_000:
                raise ValueError()
            restored=json.loads(uploaded.getvalue())
            if not isinstance(restored.get('drafts'),dict) or not isinstance(restored.get('history'),list) or not isinstance(restored.get('favorites'),list):
                raise ValueError()
            if any(tid not in by_id for tid in restored['drafts']) or any(not isinstance(h,dict) or h.get('test_id') not in by_id or not all(k in h for k in ['id','title','skill','date','responses','results','score','total']) for h in restored['history']):
                raise ValueError()
            for tid,d in restored['drafts'].items():
                if not isinstance(d,dict) or not isinstance(d.get('answers'),dict) or not isinstance(d.get('remaining'),(int,float)):
                    raise ValueError()
                if any(k not in {str(i) for i in range(1,41)} or not isinstance(v,str) or len(v)>120 for k,v in d['answers'].items()):
                    raise ValueError()
                d['deadline']=None
            if not isinstance(restored.get('highlights',{}),dict):
                raise ValueError()
            from reader import validate_annotations
            import re
            for rid,sections in restored.get('highlights',{}).items():
                if rid not in resources or not isinstance(sections,dict):
                    raise ValueError()
                for section,value in sections.items():
                    if not re.fullmatch(r'word-text|pdf-(?:region|text)-[1-9]\d*',section):
                        raise ValueError()
                    validate_annotations(value,'image' if section.startswith('pdf-region-') else 'text')
            st.session_state.progress=restored
            write_json(STATE,restored)
            for key in list(st.session_state):
                if key.startswith('ans-'):
                    del st.session_state[key]
            st.success('Đã khôi phục. Chọn Kho đề để tiếp tục.')
            st.rerun()
        except (ValueError,TypeError,KeyError):
            st.error('Bản sao lưu không hợp lệ. Chọn file JSON được xuất từ Bingo.')
    st.divider()
    st.subheader('Dữ liệu đề thi')
    st.write(f"Đã lập danh mục {len(tests)} bài. {sum(t['key_count']==40 for t in tests)} bài có khóa chấm đủ 40 câu. Các bài thiếu audio hoặc khóa chấm được ghi rõ trong kho đề.")
    st.write('Ở máy này, Bingo đọc trực tiếp hai ZIP gốc. Trên cloud, ứng dụng tải gói mã hóa của bài bạn đang chọn từ GitHub Releases. Cần cấu hình BINGO_PASSWORD, DATA_REPOSITORY, DATA_RELEASE và DATA_KEY trong Streamlit Secrets.')
    st.caption('Giao diện thích nghi từ Dreamy / Ginger Ninja của Donny Burnside theo giấy phép Attribution–ShareAlike kèm mẫu. Đề thi và audio là tài liệu do chủ ứng dụng cung cấp.')


if page=='Sao lưu & hướng dẫn':
    backup()
elif page=='Lịch sử' and st.session_state.get('show_review') and st.session_state.get('active') in by_id:
    practice(by_id[st.session_state.active])
elif page=='Lịch sử':
    history()
elif st.session_state.get('active') in by_id:
    practice(by_id[st.session_state.active])
else:
    library()

st.html('<footer class="footer">Bingo · Your personal IELTS journal<br>Dreamy adapted from <a href="http://www.ginger-ninja.net/" target="_blank" rel="noopener">Ginger Ninja! / Donny Burnside</a> · Attribution–ShareAlike</footer>')
